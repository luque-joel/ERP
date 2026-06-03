import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from datetime import datetime

def generate_sales_excel(fecha_inicio, fecha_fin, facturas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    title_font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid') # Slate Grey
    
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='D97706', end_color='D97706', fill_type='solid') # Amber
    
    data_font = Font(name='Calibri', size=11)
    bold_font = Font(name='Calibri', size=11, bold=True)
    
    border_side = Side(border_style="thin", color="D1D5DB")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Title Row
    ws.merge_cells('A1:K1')
    title_cell = ws['A1']
    title_cell.value = f"REPORTE AUTOMATIZADO DE VENTAS ({fecha_inicio} a {fecha_fin})"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40
    
    # Headers
    headers = [
        "Nº Factura", "Fecha Emisión", "Cliente", "Cédula/RUC", 
        "Subtotal 15%", "Subtotal 0%", "Descuento", "IVA (15%)", 
        "Total Neto", "Estado Factura", "Estado SRI"
    ]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[3].height = 25
    
    # Data Rows
    row_num = 4
    for f in facturas:
        ws.cell(row=row_num, column=1, value=f.numero_factura).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=f.fecha_emision.strftime("%Y-%m-%d %H:%M")).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=3, value=f.cliente.nombre).alignment = Alignment(horizontal='left')
        ws.cell(row=row_num, column=4, value=f.cliente.cedula_ruc).alignment = Alignment(horizontal='center')
        
        # Numbers
        ws.cell(row=row_num, column=5, value=float(f.subtotal_15)).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=6, value=float(f.subtotal_0)).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=7, value=float(f.total_descuento)).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=8, value=float(f.valor_iva)).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=9, value=float(f.total_pagar)).number_format = '$#,##0.00'
        
        ws.cell(row=row_num, column=10, value=f.estado).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=11, value=f.estado_sri).alignment = Alignment(horizontal='center')
        
        for col_idx in range(1, 12):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            
        row_num += 1
        
    # Totals Row
    ws.cell(row=row_num, column=3, value="TOTALES").font = bold_font
    ws.cell(row=row_num, column=3).alignment = Alignment(horizontal='right')
    
    # Formulas for sums
    if row_num > 4:
        for col_let, col_idx in [('E', 5), ('F', 6), ('G', 7), ('H', 8), ('I', 9)]:
            total_cell = ws.cell(row=row_num, column=col_idx)
            total_cell.value = f"=SUM({col_let}4:{col_let}{row_num-1})"
            total_cell.font = bold_font
            total_cell.number_format = '$#,##0.00'
            total_cell.border = thin_border
    else:
        for col_idx in [5, 6, 7, 8, 9]:
            ws.cell(row=row_num, column=col_idx, value=0.0).font = bold_font
            ws.cell(row=row_num, column=col_idx).number_format = '$#,##0.00'
            ws.cell(row=row_num, column=col_idx).border = thin_border
        
    # Double bottom border for totals
    double_bottom = Border(top=border_side, bottom=Side(border_style="double", color="000000"))
    for col_idx in range(1, 12):
        ws.cell(row=row_num, column=col_idx).border = double_bottom
        
    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        if col_letter == 'A':
            ws.column_dimensions[col_letter].width = 18
            continue
        for cell in col:
            if cell.value:
                if cell.row == 1:
                    continue
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_purchases_excel(fecha_inicio, fecha_fin, compras):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compras"
    
    ws.views.sheetView[0].showGridLines = True
    
    title_font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
    
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='D97706', end_color='D97706', fill_type='solid')
    
    data_font = Font(name='Calibri', size=11)
    bold_font = Font(name='Calibri', size=11, bold=True)
    
    border_side = Side(border_style="thin", color="D1D5DB")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Title Row
    ws.merge_cells('A1:I1')
    title_cell = ws['A1']
    title_cell.value = f"REPORTE AUTOMATIZADO DE COMPRAS A PROVEEDORES ({fecha_inicio} a {fecha_fin})"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 40
    
    # Headers
    headers = [
        "ID Compra", "Fecha Compra", "Proveedor", "RUC Proveedor", 
        "Factura Proveedor", "Subtotal", "IVA", "Total", "Estado"
    ]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[3].height = 25
    
    # Data Rows
    row_num = 4
    for c in compras:
        ws.cell(row=row_num, column=1, value=c.id).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=c.fecha_compra.strftime("%Y-%m-%d")).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=3, value=c.proveedor.nombre_empresa).alignment = Alignment(horizontal='left')
        ws.cell(row=row_num, column=4, value=c.proveedor.ruc).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=5, value=c.documento_referencia).alignment = Alignment(horizontal='center')
        
        # Numbers
        ws.cell(row=row_num, column=6, value=float(c.subtotal)).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=7, value=float(c.iva)).number_format = '$#,##0.00'
        ws.cell(row=row_num, column=8, value=float(c.total)).number_format = '$#,##0.00'
        
        ws.cell(row=row_num, column=9, value=c.estado).alignment = Alignment(horizontal='center')
        
        for col_idx in range(1, 10):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            
        row_num += 1
        
    # Totals Row
    ws.cell(row=row_num, column=3, value="TOTALES").font = bold_font
    ws.cell(row=row_num, column=3).alignment = Alignment(horizontal='right')
    
    # Formulas for sums
    if row_num > 4:
        for col_let, col_idx in [('F', 6), ('G', 7), ('H', 8)]:
            total_cell = ws.cell(row=row_num, column=col_idx)
            total_cell.value = f"=SUM({col_let}4:{col_let}{row_num-1})"
            total_cell.font = bold_font
            total_cell.number_format = '$#,##0.00'
            total_cell.border = thin_border
    else:
        for col_idx in [6, 7, 8]:
            ws.cell(row=row_num, column=col_idx, value=0.0).font = bold_font
            ws.cell(row=row_num, column=col_idx).number_format = '$#,##0.00'
            ws.cell(row=row_num, column=col_idx).border = thin_border
        
    # Double bottom border for totals
    double_bottom = Border(top=border_side, bottom=Side(border_style="double", color="000000"))
    for col_idx in range(1, 10):
        ws.cell(row=row_num, column=col_idx).border = double_bottom
        
    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                if cell.row == 1:
                    continue
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_sales_pdf(fecha_inicio, fecha_fin, facturas):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )
    
    story = []
    styles = getSampleStyleSheet()
    
    # Styles
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1F2937'),
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'SubtitleStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#4B5563'),
        spaceAfter=20
    )
    
    header_cell_style = ParagraphStyle(
        'HeaderCellStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.white,
        alignment=1 # Center
    )
    
    cell_style = ParagraphStyle(
        'CellStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor('#1F2937'),
        alignment=1 # Center
    )
    
    cell_left_style = ParagraphStyle(
        'CellLeftStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor('#1F2937'),
        alignment=0 # Left
    )
    
    cell_bold_style = ParagraphStyle(
        'CellBoldStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.HexColor('#1F2937'),
        alignment=1 # Center
    )

    cell_bold_right_style = ParagraphStyle(
        'CellBoldRightStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.HexColor('#1F2937'),
        alignment=2 # Right
    )

    # Document Header
    story.append(Paragraph("REPORTE AUTOMATIZADO DE VENTAS", title_style))
    story.append(Paragraph(f"MotoTaller ERP - Periodo: {fecha_inicio} a {fecha_fin} | Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
    
    # Table headers
    headers = [
        Paragraph("<b>Nº Factura</b>", header_cell_style),
        Paragraph("<b>Fecha</b>", header_cell_style),
        Paragraph("<b>Cliente</b>", header_cell_style),
        Paragraph("<b>Cedula/RUC</b>", header_cell_style),
        Paragraph("<b>Subtotal 15%</b>", header_cell_style),
        Paragraph("<b>Subtotal 0%</b>", header_cell_style),
        Paragraph("<b>Desc.</b>", header_cell_style),
        Paragraph("<b>IVA</b>", header_cell_style),
        Paragraph("<b>Total</b>", header_cell_style),
        Paragraph("<b>Estado SRI</b>", header_cell_style),
    ]
    
    data = [headers]
    
    subtotal_15_total = 0
    subtotal_0_total = 0
    descuento_total = 0
    iva_total = 0
    neto_total = 0
    
    for f in facturas:
        subtotal_15_total += f.subtotal_15
        subtotal_0_total += f.subtotal_0
        descuento_total += f.total_descuento
        iva_total += f.valor_iva
        neto_total += f.total_pagar
        
        row = [
            Paragraph(f.numero_factura, cell_style),
            Paragraph(f.fecha_emision.strftime("%d/%m/%Y"), cell_style),
            Paragraph(f.cliente.nombre[:22], cell_left_style),
            Paragraph(f.cliente.cedula_ruc, cell_style),
            Paragraph(f"${f.subtotal_15:.2f}", cell_style),
            Paragraph(f"${f.subtotal_0:.2f}", cell_style),
            Paragraph(f"${f.total_descuento:.2f}", cell_style),
            Paragraph(f"${f.valor_iva:.2f}", cell_style),
            Paragraph(f"${f.total_pagar:.2f}", cell_style),
            Paragraph(f.estado_sri, cell_style),
        ]
        data.append(row)
        
    # Add summary row
    summary_row = [
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph("<b>TOTALES</b>", cell_bold_right_style),
        Paragraph("", cell_style),
        Paragraph(f"<b>${subtotal_15_total:.2f}</b>", cell_bold_style),
        Paragraph(f"<b>${subtotal_0_total:.2f}</b>", cell_bold_style),
        Paragraph(f"<b>${descuento_total:.2f}</b>", cell_bold_style),
        Paragraph(f"<b>${iva_total:.2f}</b>", cell_bold_style),
        Paragraph(f"<b>${neto_total:.2f}</b>", cell_bold_style),
        Paragraph("", cell_style),
    ]
    data.append(summary_row)
    
    # Render table
    col_widths = [80, 55, 120, 65, 60, 60, 50, 50, 60, 70]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#D97706')), # Amber
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#D1D5DB')),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('TOPPADDING', (0,0), (-1,0), 6),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#F9FAFB')]),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#FEF3C7')), # Light Amber for totals
        ('LINEABOVE', (0,-1), (-1,-1), 1, colors.HexColor('#D97706')),
        ('BOTTOMPADDING', (0,-1), (-1,-1), 6),
        ('TOPPADDING', (0,-1), (-1,-1), 6),
    ]))
    
    story.append(t)
    doc.build(story)
    
    buffer.seek(0)
    return buffer.getvalue()

def generate_purchases_pdf(fecha_inicio, fecha_fin, compras):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    
    story = []
    styles = getSampleStyleSheet()
    
    # Styles
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1F2937'),
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'SubtitleStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#4B5563'),
        spaceAfter=20
    )
    
    header_cell_style = ParagraphStyle(
        'HeaderCellStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        textColor=colors.white,
        alignment=1 # Center
    )
    
    cell_style = ParagraphStyle(
        'CellStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=colors.HexColor('#1F2937'),
        alignment=1 # Center
    )
    
    cell_left_style = ParagraphStyle(
        'CellLeftStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=colors.HexColor('#1F2937'),
        alignment=0 # Left
    )
    
    cell_bold_style = ParagraphStyle(
        'CellBoldStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        textColor=colors.HexColor('#1F2937'),
        alignment=1 # Center
    )

    cell_bold_right_style = ParagraphStyle(
        'CellBoldRightStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        textColor=colors.HexColor('#1F2937'),
        alignment=2 # Right
    )

    # Document Header
    story.append(Paragraph("REPORTE AUTOMATIZADO DE COMPRAS A PROVEEDORES", title_style))
    story.append(Paragraph(f"MotoTaller ERP - Periodo: {fecha_inicio} a {fecha_fin} | Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
    
    # Table headers
    headers = [
        Paragraph("<b>ID Compra</b>", header_cell_style),
        Paragraph("<b>Fecha</b>", header_cell_style),
        Paragraph("<b>Proveedor</b>", header_cell_style),
        Paragraph("<b>RUC Proveedor</b>", header_cell_style),
        Paragraph("<b>Factura Prov.</b>", header_cell_style),
        Paragraph("<b>Subtotal</b>", header_cell_style),
        Paragraph("<b>IVA</b>", header_cell_style),
        Paragraph("<b>Total</b>", header_cell_style),
        Paragraph("<b>Estado</b>", header_cell_style),
    ]
    
    data = [headers]
    
    subtotal_total = 0
    iva_total = 0
    total_total = 0
    
    for c in compras:
        subtotal_total += c.subtotal
        iva_total += c.iva
        total_total += c.total
        
        row = [
            Paragraph(str(c.id), cell_style),
            Paragraph(c.fecha_compra.strftime("%d/%m/%Y"), cell_style),
            Paragraph(c.proveedor.nombre_empresa[:30], cell_left_style),
            Paragraph(c.proveedor.ruc, cell_style),
            Paragraph(c.documento_referencia or "S/N", cell_style),
            Paragraph(f"${c.subtotal:.2f}", cell_style),
            Paragraph(f"${c.iva:.2f}", cell_style),
            Paragraph(f"${c.total:.2f}", cell_style),
            Paragraph(c.estado, cell_style),
        ]
        data.append(row)
        
    # Add summary row
    summary_row = [
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph("<b>TOTALES</b>", cell_bold_right_style),
        Paragraph("", cell_style),
        Paragraph("", cell_style),
        Paragraph(f"<b>${subtotal_total:.2f}</b>", cell_bold_style),
        Paragraph(f"<b>${iva_total:.2f}</b>", cell_bold_style),
        Paragraph(f"<b>${total_total:.2f}</b>", cell_bold_style),
        Paragraph("", cell_style),
    ]
    data.append(summary_row)
    
    # Render table
    col_widths = [60, 65, 170, 85, 90, 70, 60, 70, 60]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#D97706')), # Amber
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#D1D5DB')),
        ('BOTTOMPADDING', (0,0), (-1,0), 6),
        ('TOPPADDING', (0,0), (-1,0), 6),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#F9FAFB')]),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#FEF3C7')), # Light Amber for totals
        ('LINEABOVE', (0,-1), (-1,-1), 1, colors.HexColor('#D97706')),
        ('BOTTOMPADDING', (0,-1), (-1,-1), 6),
        ('TOPPADDING', (0,-1), (-1,-1), 6),
    ]))
    
    story.append(t)
    doc.build(story)
    
    buffer.seek(0)
    return buffer.getvalue()
